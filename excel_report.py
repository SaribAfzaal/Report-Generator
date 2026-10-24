from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
# A couple of colors and styles we reuse throughout the report.
HEADER_FILL = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=16, bold=True, color="2E5090")
SUBTITLE_FONT = Font(size=10, italic=True, color="666666")
THIN_BORDER = Border(bottom=Side(style="thin", color="DDDDDD"))

def style_header_row(worksheet, row_number, num_columns):
    # Make a row look like a table header: bold white text on a blue background.
    for col_number in range(1, num_columns + 1):
        cell = worksheet.cell(row=row_number, column=col_number)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

def autosize_columns(worksheet, df, start_column=1):
    # Widen each column so the text isn't cut off.
    for i, column_name in enumerate(df.columns):
        column_letter = get_column_letter(start_column + i)
        longest_value = max(
            [len(str(column_name))] + [len(str(v)) for v in df[column_name]]
        )
        worksheet.column_dimensions[column_letter].width = longest_value + 4

def write_dataframe(worksheet, df, start_row):
    # Write a pandas DataFrame into the worksheet, starting at start_row.
    # Header row
    for col_number, column_name in enumerate(df.columns, start=1):
        worksheet.cell(row=start_row, column=col_number, value=column_name)
    style_header_row(worksheet, start_row, len(df.columns))
    # Data rows
    for row_offset, row_data in enumerate(df.itertuples(index=False), start=1):
        row_number = start_row + row_offset
        for col_number, value in enumerate(row_data, start=1):
            cell = worksheet.cell(row=row_number, column=col_number, value=value)
            cell.border = THIN_BORDER
            # Right-align numbers so they line up nicely
            if isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal="right")
    autosize_columns(worksheet, df)
    last_row = start_row + len(df)
    return last_row

def write_excel_report(output_path, report_title, generated_on, raw_data, summary, group_column, value_column):
    workbook = Workbook()
    # Summary sheet
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet["A1"] = report_title
    summary_sheet["A1"].font = TITLE_FONT
    summary_sheet["A2"] = f"Generated on {generated_on}"
    summary_sheet["A2"].font = SUBTITLE_FONT
    # A couple of quick headline numbers above the table
    summary_sheet["A4"] = "Total rows in source data:"
    summary_sheet["B4"] = len(raw_data)
    summary_sheet["A5"] = f"Overall total ({value_column}):"
    summary_sheet["B5"] = round(raw_data[value_column].sum(), 2)
    summary_sheet["A6"] = f"Overall average ({value_column}):"
    summary_sheet["B6"] = round(raw_data[value_column].mean(), 2)
    for row in (4, 5, 6):
        summary_sheet[f"A{row}"].font = Font(bold=True)
    write_dataframe(summary_sheet, summary, start_row=8)
    # Raw data sheet
    data_sheet = workbook.create_sheet("Raw Data")
    data_sheet["A1"] = "Raw data"
    data_sheet["A1"].font = TITLE_FONT
    write_dataframe(data_sheet, raw_data, start_row=3)
    workbook.save(output_path)